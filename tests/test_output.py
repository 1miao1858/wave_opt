from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.config import (
    Config,
    SubProblemConfig,
    SweepConfig,
    MipConfig,
    ValidationConfig,
)
from src.data_loader import load_orders, load_inventory_snapshots
from src.sweep import SweepRunner, SweepResult, WindowResult
from src.output import write_excel


def _make_config(N_max=2):
    return Config(
        sub_problems={
            "fei_jia_gong_le_20": SubProblemConfig(
                wave_type="非加工",
                size_class="le_20",
                N_max=N_max,
            ),
        },
        sweep=SweepConfig(
            num_pickers=4,
            pick_time_per_order=120,
            num_machines=2,
            process_time_per_jian=36,
            K_ws=8,
            daily_cutoff="14:00",
            daily_deadline="18:00",
            picker_shift_start="08:00",
            picker_shift_end="18:00",
            machine_shift_start="08:00",
            machine_shift_end="18:00",
        ),
        candidate_W=["1h"],
        mip=MipConfig(
            solver="gurobi",
            time_limit=60,
            mip_gap=0.05,
            fallback="greedy",
        ),
        validation=ValidationConfig(
            recomputed_hit_rate=True,
            boundary_check=True,
            inventory_consistency_sample=5,
            stockout_filter_sample=10,
        ),
    )


FIXTURE = Path(__file__).parent / "fixtures" / "tiny_case"


def _build_sweep_results():
    orders = load_orders(FIXTURE / "orders.csv")
    snaps = load_inventory_snapshots(FIXTURE / "inventory_snapshots.csv")
    cfg = _make_config(N_max=2)
    runner = SweepRunner(cfg=cfg, all_orders=orders, snapshots=snaps)
    return {sp: runner.run_sub_problem(sp) for sp in cfg.sub_problems}


def test_excel_sheet1_has_summary(tmp_path):
    results = _build_sweep_results()
    out_file = tmp_path / "out.xlsx"
    write_excel(results, out_file)

    wb = load_workbook(out_file)
    assert "Sheet1_汇总" in wb.sheetnames
    ws = wb["Sheet1_汇总"]
    # 表头应有:子问题 | W | 总货架访问数 | 平均命中率 | 命中率方差 | 平均子波规模 | 拣货利用率 | 加工利用率 | 截单可行 | 备注
    headers = [c.value for c in ws[1]]
    assert "子问题" in headers
    assert "总货架访问数" in headers
    assert "平均命中率" in headers
    # 数据行:1 子问题 × 1 W = 1 行
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) >= 1
    row = data_rows[0]
    assert row[0] == "fei_jia_gong_le_20"
    assert isinstance(row[2], (int, float))  # 总货架访问数
    assert isinstance(row[3], (int, float))  # 平均命中率
