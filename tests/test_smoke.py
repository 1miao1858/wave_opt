"""端到端 smoke:小样本能跑通,5 Sheet 齐全,最优解对齐手算。"""
from pathlib import Path
import json
import subprocess
import sys

import pytest
from openpyxl import load_workbook

PROJECT = Path(__file__).parent.parent
FIXTURE = PROJECT / "tests" / "fixtures" / "tiny_case"


def test_smoke_run_full_pipeline(tmp_path):
    """跑 tiny_case 端到端,验证 5 Sheet + 最优解对齐。"""
    config_text = """
sub_problems:
  fei_jia_gong_le_20:
    wave_type: 非加工
    size_class: le_20
    N_max: 2
sweep:
  num_pickers: 4
  pick_time_per_order: 120
  num_machines: 2
  process_time_per_jian: 36
  K_ws: 8
  daily_cutoff: "14:00"
  daily_deadline: "18:00"
  picker_shift_start: "08:00"
  picker_shift_end: "18:00"
  machine_shift_start: "08:00"
  machine_shift_end: "18:00"
candidate_W: ["1h"]
mip:
  solver: gurobi
  time_limit: 60
  mip_gap: 0.05
  fallback: greedy
validation:
  recomputed_hit_rate: true
  boundary_check: true
  inventory_consistency_sample: 5
  stockout_filter_sample: 10
"""
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text(config_text, encoding="utf-8")
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "orders.csv").write_text(
        (FIXTURE / "orders.csv").read_text(encoding="utf-8"), encoding="utf-8"
    )
    (tmp_path / "data" / "inventory_snapshots.csv").write_text(
        (FIXTURE / "inventory_snapshots.csv").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "src.main",
            "--config",
            str(cfg_path),
            "--data-dir",
            str(tmp_path / "data"),
            "--output",
            str(tmp_path / "out.xlsx"),
        ],
        cwd=PROJECT,
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert result.returncode == 0, f"stderr: {result.stderr}"

    wb = load_workbook(tmp_path / "out.xlsx")
    expected_sheets = [
        "Sheet1_汇总",
        "Sheet2_总访问数曲线",
        "Sheet3_命中率利用率",
        "Sheet4_波次明细",
        "Sheet5_异常诊断",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"缺 sheet:{s}"

    # 最优解对齐:Sheet 1 应有总访问数 = 3
    ws1 = wb["Sheet1_汇总"]
    rows = list(ws1.iter_rows(min_row=2, values_only=True))
    assert any(r[2] == 3 for r in rows), f"应得到最优访问数 3,实际行:{rows}"

    # Sheet 4 应有 2 个子波(3 单 / N_max=2)
    ws4 = wb["Sheet4_波次明细"]
    wave_rows = list(ws4.iter_rows(min_row=2, values_only=True))
    assert len(wave_rows) >= 2
