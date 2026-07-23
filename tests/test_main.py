import subprocess
import sys
from pathlib import Path

import pytest


def test_main_runs_tiny_case(tmp_path):
    """端到端跑通 tiny_case:生成 xlsx 输出。"""
    project_root = Path(__file__).parent.parent
    fixture = project_root / "tests" / "fixtures" / "tiny_case"

    # 临时 config:tiny_case 的 3 单 N_max=2
    config_text = """
sub_problems:
  fei_jia_gong_le_20:
    wave_type: 非加工
    size_class: le_20
    N_max: 2
sweep:
  num_pickers: 4
  pick_time_per_order: 120
  num_machines: 2
  process_time_per_jian: 36
  K_ws: 8
  daily_cutoff: "14:00"
  daily_deadline: "18:00"
  picker_shift_start: "08:00"
  picker_shift_end: "18:00"
  machine_shift_start: "08:00"
  machine_shift_end: "18:00"
candidate_W: ["1h"]
mip:
  solver: gurobi
  time_limit: 60
  mip_gap: 0.05
  fallback: greedy
validation:
  recomputed_hit_rate: true
  boundary_check: true
  inventory_consistency_sample: 5
  stockout_filter_sample: 10
"""
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text(config_text, encoding="utf-8")

    # 复制 tiny_case 数据到 tmp_path/data
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "orders.csv").write_text(
        (fixture / "orders.csv").read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    (tmp_path / "data" / "inventory_snapshots.csv").write_text(
        (fixture / "inventory_snapshots.csv").read_text(encoding="utf-8"),
        encoding="utf-8",
    )

    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "src.main",
            "--config",
            str(cfg_path),
            "--data-dir",
            str(tmp_path / "data"),
            "--output",
            str(tmp_path / "out.xlsx"),
        ],
        cwd=project_root,
        capture_output=True,
        text=True,
        timeout=120,
    )

    assert result.returncode == 0, f"stderr: {result.stderr}"
    out_file = tmp_path / "out.xlsx"
    assert out_file.exists()
    # 应有 5 个 sheet
    from openpyxl import load_workbook

    wb = load_workbook(out_file)
    assert "Sheet1_汇总" in wb.sheetnames
    assert "Sheet4_波次明细" in wb.sheetnames
