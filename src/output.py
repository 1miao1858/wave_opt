"""Excel 5 Sheet 输出。"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

from src.sweep import SweepResult, WindowResult

BEST_W_FILL = PatternFill("solid", fgColor="FFD700")  # 金色高亮
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _aggregate_sweep_metrics(result: SweepResult):
    """聚合 SweepResult -> 每 W 的汇总行。"""
    rows = []
    for W, window_results in result.window_results_by_W.items():
        if not window_results:
            continue
        feasible_windows = [r for r in window_results if r.feasible]
        if not feasible_windows:
            rows.append({
                "子问题": result.sub_problem_key,
                "W": W,
                "总货架访问数": 0,
                "平均命中率": 0.0,
                "命中率方差": 0.0,
                "平均子波规模": 0.0,
                "拣货利用率": 0.0,
                "加工利用率": None,
                "截单可行": False,
                "备注": ";".join(window_results[0].notes) if window_results else "",
            })
            continue
        total_visits = sum(r.total_visits for r in feasible_windows)
        avg_hit = (
            sum(r.solution.hit_rate for r in feasible_windows)
            / len(feasible_windows)
        )
        # 命中率方差
        if len(feasible_windows) > 1:
            mean = avg_hit
            var = sum(
                (r.solution.hit_rate - mean) ** 2 for r in feasible_windows
            ) / len(feasible_windows)
        else:
            var = 0.0
        avg_subwave_size = (
            sum(
                len(w.orders)
                for r in feasible_windows
                for w in r.solution.wave_assignments
            )
            / max(
                1,
                sum(len(r.solution.wave_assignments) for r in feasible_windows),
            )
        )
        avg_picker = (
            sum(r.picker_utilization for r in feasible_windows)
            / len(feasible_windows)
        )
        avg_machine = None
        machine_vals = [
            r.machine_utilization
            for r in feasible_windows
            if r.machine_utilization is not None
        ]
        if machine_vals:
            avg_machine = sum(machine_vals) / len(machine_vals)
        rows.append({
            "子问题": result.sub_problem_key,
            "W": W,
            "总货架访问数": total_visits,
            "平均命中率": round(avg_hit, 4),
            "命中率方差": round(var, 6),
            "平均子波规模": round(avg_subwave_size, 2),
            "拣货利用率": round(avg_picker, 4),
            "加工利用率": round(avg_machine, 4) if avg_machine is not None else None,
            "截单可行": all(r.feasible for r in window_results),
            "备注": "",
        })
    return rows


def _write_sheet1(wb: Workbook, results: dict[str, SweepResult]) -> None:
    ws = wb.create_sheet("Sheet1_汇总")
    headers = [
        "子问题",
        "W",
        "总货架访问数",
        "平均命中率",
        "命中率方差",
        "平均子波规模",
        "拣货利用率",
        "加工利用率",
        "截单可行",
        "备注",
    ]
    ws.append(headers)
    for col in ws[1]:
        col.fill = HEADER_FILL
        col.font = HEADER_FONT

    # 标记每个子问题最优 W(总访问数最小、截单可行)
    best_W_per_sub = {}
    for sp, result in results.items():
        feasible = {
            W: sum(r.total_visits for r in rs if r.feasible)
            for W, rs in result.window_results_by_W.items()
            if rs and all(r.feasible for r in rs)
        }
        if feasible:
            best_W_per_sub[sp] = min(feasible, key=feasible.get)

    row_idx = 2
    for sp, result in results.items():
        rows = _aggregate_sweep_metrics(result)
        for r in rows:
            ws.append([
                r["子问题"],
                r["W"],
                r["总货架访问数"],
                r["平均命中率"],
                r["命中率方差"],
                r["平均子波规模"],
                r["拣货利用率"],
                r["加工利用率"],
                "✓" if r["截单可行"] else "✗",
                r["备注"],
            ])
            if best_W_per_sub.get(sp) == r["W"]:
                for c in ws[row_idx]:
                    c.fill = BEST_W_FILL
            row_idx += 1


def _write_sheet2(wb: Workbook, results: dict[str, SweepResult]) -> None:
    ws = wb.create_sheet("Sheet2_总访问数曲线")
    headers = ["子问题", "W", "全日总货架访问数"]
    ws.append(headers)
    for col in ws[1]:
        col.fill = HEADER_FILL
        col.font = HEADER_FONT

    for sp, result in results.items():
        for W, window_results in result.window_results_by_W.items():
            if not window_results:
                continue
            total = sum(r.total_visits for r in window_results if r.feasible)
            ws.append([sp, W, total])


def _write_sheet3(wb: Workbook, results: dict[str, SweepResult]) -> None:
    ws = wb.create_sheet("Sheet3_命中率利用率")
    headers = ["子问题", "W", "平均命中率", "拣货利用率", "加工利用率"]
    ws.append(headers)
    for col in ws[1]:
        col.fill = HEADER_FILL
        col.font = HEADER_FONT

    for sp, result in results.items():
        for W, window_results in result.window_results_by_W.items():
            if not window_results:
                continue
            feasible = [r for r in window_results if r.feasible]
            if not feasible:
                continue
            avg_hit = sum(r.solution.hit_rate for r in feasible) / len(feasible)
            avg_picker = sum(r.picker_utilization for r in feasible) / len(feasible)
            machine_vals = [
                r.machine_utilization
                for r in feasible
                if r.machine_utilization is not None
            ]
            avg_machine = (
                sum(machine_vals) / len(machine_vals) if machine_vals else None
            )
            ws.append([
                sp,
                W,
                round(avg_hit, 4),
                round(avg_picker, 4),
                round(avg_machine, 4) if avg_machine is not None else None,
            ])


def _write_sheet4(wb: Workbook, results: dict[str, SweepResult]) -> None:
    ws = wb.create_sheet("Sheet4_波次明细")
    headers = [
        "波次ID",
        "子问题",
        "窗口ID",
        "子波序号",
        "触发时刻",
        "窗口时段",
        "订单数",
        "订单列表",
        "访问货架数",
        "hits",
        "命中率",
        "库存消耗明细",
    ]
    ws.append(headers)
    for col in ws[1]:
        col.fill = HEADER_FILL
        col.font = HEADER_FONT

    for sp, result in results.items():
        if not result.best_W:
            continue
        # 只输出最优 W 下的窗口
        best_W = result.best_W
        for w_idx, w_result in enumerate(
            result.window_results_by_W.get(best_W, [])
        ):
            if not w_result.feasible or w_result.solution is None:
                continue
            sol = w_result.solution
            window_id = (
                w_result.window_start.isoformat()
                if w_result.window_start
                else f"W{w_idx}"
            )
            for sw in sol.wave_assignments:
                # hits = Σ_{s} |{k picked}|
                shelf_sku_hits: dict[str, set[str]] = {}
                for (sku, shelf), q in sw.pick_qty.items():
                    if q > 0:
                        shelf_sku_hits.setdefault(shelf, set()).add(sku)
                hits = sum(len(s) for s in shelf_sku_hits.values())
                visits = len(sw.visited_shelves)
                hit_rate = (hits / visits) if visits > 0 else 0.0
                consumption_detail = {
                    sku: {shelf: q}
                    for (sku, shelf), q in sw.pick_qty.items()
                    if q > 0
                }
                ws.append([
                    f"{sp}-{w_idx}-{sw.subwave_idx}",
                    sp,
                    window_id,
                    sw.subwave_idx,
                    (
                        w_result.window_start.isoformat()
                        if w_result.window_start
                        else ""
                    ),
                    window_id,
                    len(sw.orders),
                    ",".join(o.order_id for o in sw.orders),
                    visits,
                    hits,
                    round(hit_rate, 4),
                    json.dumps(consumption_detail, ensure_ascii=False),
                ])


def _write_sheet5(wb: Workbook, results: dict[str, SweepResult]) -> None:
    ws = wb.create_sheet("Sheet5_异常诊断")
    headers = ["类别", "子问题", "窗口", "详情"]
    ws.append(headers)
    for col in ws[1]:
        col.fill = HEADER_FILL
        col.font = HEADER_FONT

    for sp, result in results.items():
        for W, window_results in result.window_results_by_W.items():
            for w_idx, w in enumerate(window_results):
                window_id = (
                    w.window_start.isoformat() if w.window_start else f"W{w_idx}"
                )
                # 缺货剔除
                for oid, reason in zip(w.rejected_order_ids, w.reject_reasons):
                    ws.append([
                        "缺货剔除",
                        sp,
                        window_id,
                        f"{oid}: {reason}",
                    ])
                # 不可行窗口
                if not w.feasible:
                    ws.append([
                        "不可行窗口",
                        sp,
                        window_id,
                        f"notes={list(w.notes)}",
                    ])
                # fallback 贪心
                if w.solution and "fallback_greedy" in str(w.notes):
                    ws.append([
                        "MIP超时fallback贪心",
                        sp,
                        window_id,
                        f"notes={list(w.notes)}",
                    ])
                # 子波拆分数 > 3
                if w.solution and len(w.solution.wave_assignments) > 3:
                    ws.append([
                        "子波拆分数过多",
                        sp,
                        window_id,
                        f"子波数={len(w.solution.wave_assignments)}"
                        f"(N_max 可能过小)",
                    ])
                # MIP gap > 5%
                if (
                    w.solution
                    and w.solution.mip_gap
                    and w.solution.mip_gap > 0.05
                ):
                    ws.append([
                        "MIP gap 超阈",
                        sp,
                        window_id,
                        f"gap={w.solution.mip_gap:.4f}",
                    ])


def write_excel(results: dict[str, SweepResult], out_path: Path | str) -> None:
    """写完整 Excel(5 Sheet)。"""
    out_path = Path(out_path)
    wb = Workbook()
    wb.remove(wb.active)  # 删默认 Sheet
    _write_sheet1(wb, results)
    _write_sheet2(wb, results)
    _write_sheet3(wb, results)
    _write_sheet4(wb, results)
    _write_sheet5(wb, results)
    wb.save(out_path)
